# app.py
# Batch DOCX/TXT -> TTML (no timecodes)
# Run:
#   pip install streamlit python-docx
#   streamlit run app.py

import re
import html
import io
import zipfile
from typing import List
import streamlit as st


# -------------------- Input helpers --------------------

def decode_bytes(data: bytes) -> str:
    for enc in ("utf-8", "utf-8-sig", "windows-1252", "iso-8859-1", "macroman"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            pass
    return data.decode("utf-8", errors="replace")


def docx_bytes_to_text(data: bytes) -> str:
    from io import BytesIO
    from docx import Document  # python-docx

    doc = Document(BytesIO(data))
    lines: List[str] = []

    # Paragraphs
    for p in doc.paragraphs:
        lines.append(p.text)

    # Tables (optional but useful)
    for table in doc.tables:
        for row in table.rows:
            cells = []
            for cell in row.cells:
                cell_text = re.sub(r"\s+", " ", cell.text).strip()
                cells.append(cell_text)
            lines.append("\t".join(cells))

    return "\n".join(lines)


# -------------------- Converter --------------------

def fmt_time(t: float) -> str:
    h = int(t // 3600); t -= h * 3600
    m = int(t // 60);   t -= m * 60
    s = int(t)
    ms = int(round((t - s) * 1000))
    if ms == 1000:
        s += 1
        ms = 0
    return f"{h:02d}:{m:02d}:{s:02d}.{ms:03d}"


def word_count(text: str) -> int:
    return len(re.findall(r"\b[\w’']+\b", text, flags=re.UNICODE))


def char_count(text: str) -> int:
    t = re.sub(r"\s+", " ", text).strip()
    return len(t)


def estimate_duration(text: str, *, rate_mode: str, wpm: float, cps: float, max_dur: float) -> float:
    if rate_mode.upper() == "CPS":
        n = char_count(text)
        dur = (n / cps) if n else 1.0
    else:
        n = word_count(text)
        dur = (n / wpm) * 60.0 if n else 1.0
    return max(1.0, min(dur, max_dur))


def split_units(raw: str, split_mode: str) -> List[str]:
    raw = raw.replace("\r\n", "\n").replace("\r", "\n")
    if split_mode == "each_line":
        return [ln.strip() for ln in raw.split("\n") if ln.strip()]
    # blank_lines
    return [p.strip() for p in re.split(r"\n\s*\n+", raw) if p.strip()]


def txt_to_ttml(
    raw: str,
    *,
    split_mode: str = "blank_lines",  # blank_lines | each_line
    rate_mode: str = "WPM",           # WPM | CPS
    wpm: float = 160.0,
    cps: float = 17.0,
    gap: float = 1.0,
    lang: str = "fr",
    max_dur: float = 30.0,
) -> str:
    units = split_units(raw, split_mode)

    t = 0.0
    body = []

    for i, unit in enumerate(units, start=1):
        text = re.sub(r"\s+", " ", unit).strip()
        dur = estimate_duration(text, rate_mode=rate_mode, wpm=wpm, cps=cps, max_dur=max_dur)

        begin = t
        end = t + dur
        t = end + gap

        body.append(
            f'      <p xml:id="p{i}" begin="{fmt_time(begin)}" end="{fmt_time(end)}">{html.escape(text)}</p>'
        )

    return f'''<?xml version="1.0" encoding="UTF-8"?>
<tt xmlns="http://www.w3.org/ns/ttml"
    xmlns:tts="http://www.w3.org/ns/ttml#styling"
    xmlns:ttm="http://www.w3.org/ns/ttml#metadata"
    xml:lang="{html.escape(lang)}">
  <head>
    <styling>
      <style xml:id="s1" tts:fontFamily="sansSerif" tts:fontSize="100%" tts:textAlign="center"/>
    </styling>
    <layout>
      <region xml:id="bottom" tts:origin="10% 80%" tts:extent="80% 15%" tts:displayAlign="after"/>
    </layout>
  </head>
  <body region="bottom" style="s1">
    <div>
{chr(10).join(body)}
    </div>
  </body>
</tt>
'''


def base_name(filename: str) -> str:
    return filename.rsplit(".", 1)[0]


# -------------------- Streamlit UI (EN) --------------------

st.set_page_config(page_title="Batch DOCX/TXT → TTML", layout="centered")
st.title("Batch DOCX/TXT → TTML converter (no timecodes)")

uploads = st.file_uploader(
    "Upload one or more .txt / .docx files",
    type=["txt", "docx"],
    accept_multiple_files=True,
)

split_label = st.selectbox(
    "Cue splitting",
    ["Blank lines (1 paragraph = 1 cue)", "Each line = 1 cue"],
)

rate_label = st.selectbox(
    "Speech rate mode",
    ["CPS (characters per second)", "WPM (words per minute)"],
)

col1, col2 = st.columns(2)
with col1:
    if rate_label.startswith("WPM"):
        rate_value = st.number_input("WPM value", min_value=60, max_value=300, value=160, step=5)
    else:
        rate_value = st.number_input("CPS value", min_value=5.0, max_value=30.0, value=17.0, step=0.5)

with col2:
    gap = st.number_input("Gap between cues (seconds)", min_value=0.2, max_value=5.0, value=1.0, step=0.1)

lang = st.text_input("Language (xml:lang)", value="fr")
max_dur = st.number_input("Max cue duration (seconds)", min_value=5.0, max_value=60.0, value=30.0, step=1.0)

if uploads:
    split_mode = "blank_lines" if split_label.startswith("Blank lines") else "each_line"
    rate_mode = "WPM" if rate_label.startswith("WPM") else "CPS"

    st.caption(f"Files selected: {len(uploads)}")

    if st.button("Generate TTML (ZIP)"):
        zip_buf = io.BytesIO()

        with zipfile.ZipFile(zip_buf, "w", compression=zipfile.ZIP_DEFLATED) as zf:
            for up in uploads:
                name = up.name
                data = up.getvalue()

                # read as text
                if name.lower().endswith(".docx"):
                    raw = docx_bytes_to_text(data)
                else:
                    raw = decode_bytes(data)

                ttml = txt_to_ttml(
                    raw,
                    split_mode=split_mode,
                    rate_mode=rate_mode,
                    wpm=float(rate_value),
                    cps=float(rate_value),
                    gap=float(gap),
                    lang=lang,
                    max_dur=float(max_dur),
                )

                out_name = base_name(name) + ".ttml"
                zf.writestr(out_name, ttml.encode("utf-8"))

        zip_buf.seek(0)
        st.download_button(
            "Download ZIP",
            data=zip_buf.getvalue(),
            file_name="ttml_batch.zip",
            mime="application/zip",
        )